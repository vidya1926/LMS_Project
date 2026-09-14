from pptx import Presentation
import pandas as pd

prs = Presentation("Lecture.pptx")

rows = []

# 2. Loop through slides
for slide_index, slide in enumerate(prs.slides):
    texts = []
    # 3. Loop through shapes on each slide
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            if shape.text:
                texts.append(shape.text)
    raw_text = "\n".join(texts)

    rows.append({
        "lecture_id": "Lecture_6",
        "slide_index": slide_index,
        "raw_text": raw_text
    })

# 4. Save as CSV
df = pd.DataFrame(rows)
df.to_csv("lecture_slides.csv", index=False)
print("Saved lecture_slides.csv")
